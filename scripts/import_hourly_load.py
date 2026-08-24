#!/usr/bin/env python3
"""Convert the FIIS hourly-load workbook into an API seed CSV."""

import argparse
import csv
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


HEADER_ALIASES = {
    "CODIGO": "course",
    "CÓDIGO": "course",
    "NOMBRE DEL CURSO": "course_name",
    "SECCION": "section",
    "SECCIÓN": "section",
    "SISTEMA DE EVALUACION": "evaluation",
    "SISTEMA DE EVALUACIÓN": "evaluation",
    "APELLIDOS Y NOMBRES DEL DOCENTE": "teacher",
    "APELLIDOS Y NOMBRES DE DOCENTE": "teacher",
    "TIPO": "session_type",
    "TIPO CLASE": "session_type",
    "AULA": "classroom",
    "DIA": "day",
    "DÍA": "day",
    "HORA INICIO": "start",
    "HORA FINAL": "end",
    "DNI": "teacher_dni",
    "VACANTES": "vacancies",
}

REQUIRED_TARGETS = {
    "course",
    "course_name",
    "section",
    "evaluation",
    "teacher",
    "session_type",
    "classroom",
    "day",
    "start",
    "end",
    "vacancies",
}

OUTPUT_HEADERS = [
    "codigo_facultad",
    "codigo_curso",
    "nombre_curso",
    "seccion",
    "evaluacion",
    "vacantes",
    "inicio",
    "fin",
    "aula",
    "dni_docente",
    "nombre_docente",
    "tipo",
    "dia",
]


import datetime


def clean(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime.timedelta):
        if value.seconds == 0 and value.days > 0:
            return str(value.days)
        total_hours = value.days * 24 + value.seconds // 3600
        minutes = (value.seconds % 3600) // 60
        if minutes == 0:
            return str(total_hours)
        return f"{total_hours:02d}:{minutes:02d}"
    if isinstance(value, datetime.time):
        return value.strftime("%H:%M")
    if isinstance(value, datetime.datetime):
        return value.strftime("%H:%M")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def clean_teacher_name(value: str) -> str:
    return " ".join(value.replace("\u00a0", " ").split())


def normalized_identity(value: str) -> str:
    decomposed = unicodedata.normalize("NFKD", value)
    without_marks = "".join(character for character in decomposed if not unicodedata.combining(character))
    return " ".join(re.sub(r"[^A-Z0-9]+", " ", without_marks.upper()).split())


def disambiguate_teacher_dni(rows: list[dict[str, str]]) -> int:
    names_by_dni: dict[str, list[str]] = defaultdict(list)
    for row in rows:
        row["nombre_docente"] = clean_teacher_name(row["nombre_docente"])
        if row["dni_docente"]:
            names_by_dni[row["dni_docente"]].append(row["nombre_docente"])

    ambiguous = {
        dni
        for dni, names in names_by_dni.items()
        if len({normalized_identity(name) for name in names}) > 1
    }
    canonical_names = {
        dni: Counter(names).most_common(1)[0][0]
        for dni, names in names_by_dni.items()
        if dni not in ambiguous
    }

    for row in rows:
        dni = row["dni_docente"]
        if dni in ambiguous:
            row["dni_docente"] = ""
        elif dni:
            row["nombre_docente"] = canonical_names[dni]
    return len(ambiguous)


def find_header(sheet) -> tuple[int, dict[str, int]]:
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        positions = {}
        for index, value in enumerate(row):
            if value is None:
                continue
            norm = " ".join(clean(value).upper().split())
            if norm in HEADER_ALIASES:
                positions[HEADER_ALIASES[norm]] = index
        if REQUIRED_TARGETS <= positions.keys():
            return row_number, positions
    raise ValueError("No se encontró la cabecera esperada de la carga horaria")


def previous_fallbacks(path: Path | None) -> dict[tuple[str, ...], tuple[str, str]]:
    if path is None:
        return {}
    with path.open(encoding="utf-8", newline="") as source:
        rows = csv.DictReader(source)
        return {
            (
                row["codigo_curso"].strip(),
                row["seccion"].strip(),
                row["dia"].strip(),
                row["inicio"].strip(),
                row["fin"].strip(),
            ): (row["nombre_docente"].strip(), row["tipo"].strip())
            for row in rows
        }


def convert(args: argparse.Namespace) -> int:
    workbook = load_workbook(args.workbook, data_only=True, read_only=True)
    sheet = workbook[args.sheet] if args.sheet and args.sheet in workbook.sheetnames else workbook.worksheets[0]
    header_row, columns = find_header(sheet)
    fallbacks = previous_fallbacks(args.previous)
    output: list[dict[str, str]] = []

    required = ("course", "course_name", "section", "day", "start", "end")
    for source_row_number, values in enumerate(
        sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1
    ):
        row = {
            name: clean(values[index] if index is not None and index < len(values) else None)
            for name, index in columns.items()
        }
        if "teacher_dni" not in row:
            row["teacher_dni"] = ""
        if not any(row.values()):
            continue
        missing = [name for name in required if not row[name]]
        if missing:
            print(
                f"Aviso fila {source_row_number}: omitiendo {row.get('course')} sec {row.get('section')} por campos faltantes: {', '.join(missing)}"
            )
            continue

        key = (row["course"], row["section"], row["day"], row["start"], row["end"])
        previous_teacher, previous_type = fallbacks.get(key, ("", ""))
        teacher = row["teacher"] or previous_teacher
        session_type = row["session_type"] or previous_type
        if not teacher:
            teacher = "AYUDANTE CATEDRA" if row["teacher_dni"].upper() == "SOLO PARA PC" else "NN"
        if not session_type and row["teacher_dni"].upper() == "SOLO PARA PC":
            session_type = "PC"

        output.append(
            {
                "codigo_facultad": args.faculty,
                "codigo_curso": row["course"].replace("-", ""),
                "nombre_curso": row["course_name"],
                "seccion": row["section"],
                "evaluacion": row["evaluation"],
                "vacantes": row["vacancies"] if row["vacancies"].strip().isdigit() else "0",
                "inicio": row["start"],
                "fin": row["end"],
                "aula": row["classroom"],
                "dni_docente": "" if row["teacher_dni"].upper() == "SOLO PARA PC" else row["teacher_dni"],
                "nombre_docente": teacher,
                "tipo": session_type,
                "dia": row["day"],
            }
        )

    ambiguous_dni = disambiguate_teacher_dni(output)
    output.sort(key=lambda row: (row["codigo_curso"], row["seccion"], row["dia"], row["inicio"], row["tipo"]))
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8", newline="") as target:
        writer = csv.DictWriter(target, fieldnames=OUTPUT_HEADERS, lineterminator="\n")
        writer.writeheader()
        writer.writerows(output)
    print(f"generated: {args.output} ({len(output)} rows, {ambiguous_dni} ambiguous DNI omitted)")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--sheet", default="CARGA HORARIA 2026-2")
    parser.add_argument("--faculty", default="I")
    parser.add_argument("--previous", type=Path)
    return convert(parser.parse_args())


if __name__ == "__main__":
    raise SystemExit(main())
